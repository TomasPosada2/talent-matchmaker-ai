"""
Genera una bandeja de entrada simulada en data/inbox_simulado/ con:
  - 20 correos con CV válido adjunto (.docx), en varias plantillas/formatos
  - 2 correos sin adjunto (HU-01: deben ignorarse)
  - 2 correos con adjunto en formato no soportado (HU-01: deben descartarse)
  - 1 correo con adjunto .docx vacío (HU-08: caso borde, error controlado)

No representa datos reales de personas: los CVs son sintéticos generados con Faker.
"""

import json
import random
from pathlib import Path

from docx import Document
from faker import Faker

random.seed(42)
fake = Faker("es_CO")
Faker.seed(42)

BASE_DIR = Path(__file__).resolve().parent.parent
INBOX_DIR = BASE_DIR / "data" / "inbox_simulado"

HABILIDADES_POOL = [
    "Python", "SQL", "Excel avanzado", "Power BI", "R", "Java",
    "Comunicación efectiva", "Trabajo en equipo", "AWS", "Docker",
    "Análisis de datos", "Gestión de proyectos", "Inglés B2",
]

CARGOS = [
    "Analista de Datos", "Desarrollador Backend", "Ingeniero de Software",
    "Asistente Administrativo", "Coordinador de Proyectos", "Diseñador UX",
]

UNIVERSIDADES = [
    "Universidad EIA", "Universidad de Antioquia", "Universidad Nacional",
    "Universidad Pontificia Bolivariana", "EAFIT",
]


def _crear_cv_docx(ruta: Path, plantilla: str, incluir_habilidades=True, incluir_educacion=True):
    doc = Document()
    nombre = fake.name()
    email = fake.email()
    telefono = fake.phone_number()

    if plantilla == "clasica":
        doc.add_paragraph(nombre)
        doc.add_paragraph(f"{email} | {telefono}")
        doc.add_paragraph("")
        if incluir_educacion:
            doc.add_paragraph("Educación")
            doc.add_paragraph(f"{random.choice(UNIVERSIDADES)} — {random.choice(['Ingeniería', 'Administración', 'Ciencias'])}")
        doc.add_paragraph("")
        doc.add_paragraph("Experiencia")
        doc.add_paragraph(f"{random.choice(CARGOS)} en {fake.company()} ({fake.year()}-{fake.year()})")
        doc.add_paragraph("")
        if incluir_habilidades:
            doc.add_paragraph("Habilidades")
            doc.add_paragraph(", ".join(random.sample(HABILIDADES_POOL, 4)))

    elif plantilla == "en_ingles":
        doc.add_paragraph(nombre)
        doc.add_paragraph(f"{email} / {telefono}")
        if incluir_educacion:
            doc.add_paragraph("Education")
            doc.add_paragraph(f"{random.choice(UNIVERSIDADES)} — BSc")
        doc.add_paragraph("Experience")
        doc.add_paragraph(f"{random.choice(CARGOS)} at {fake.company()}")
        if incluir_habilidades:
            doc.add_paragraph("Skills")
            doc.add_paragraph(", ".join(random.sample(HABILIDADES_POOL, 3)))

    elif plantilla == "minimalista":
        # Sin sección de educación (para probar campo faltante -> None)
        doc.add_paragraph(nombre)
        doc.add_paragraph(email)
        doc.add_paragraph("Experiencia")
        doc.add_paragraph(f"{random.choice(CARGOS)} — {fake.company()}")
        if incluir_habilidades:
            doc.add_paragraph("Habilidades")
            doc.add_paragraph(", ".join(random.sample(HABILIDADES_POOL, 5)))

    elif plantilla == "con_tabla":
        doc.add_paragraph(nombre)
        doc.add_paragraph(f"{email} | {telefono}")
        doc.add_paragraph("Educación")
        doc.add_paragraph(f"{random.choice(UNIVERSIDADES)}")
        doc.add_paragraph("Experiencia")
        tabla = doc.add_table(rows=1, cols=2)
        celdas = tabla.rows[0].cells
        celdas[0].text = random.choice(CARGOS)
        celdas[1].text = fake.company()
        doc.add_paragraph("Habilidades")
        doc.add_paragraph(", ".join(random.sample(HABILIDADES_POOL, 4)))

    doc.save(ruta)
    return nombre


def generar_bandeja():
    if INBOX_DIR.exists():
        import shutil
        shutil.rmtree(INBOX_DIR)
    INBOX_DIR.mkdir(parents=True)

    plantillas = ["clasica", "en_ingles", "minimalista", "con_tabla"]
    contador = 1

    # --- 20 correos con CV válido ---
    for i in range(20):
        email_id = f"email_{contador:03d}"
        carpeta = INBOX_DIR / email_id
        carpeta.mkdir()
        plantilla = plantillas[i % len(plantillas)]
        nombre_archivo = f"cv_{i+1}.docx"
        nombre_candidato = _crear_cv_docx(
            carpeta / nombre_archivo,
            plantilla,
            incluir_habilidades=random.random() > 0.15,
            incluir_educacion=random.random() > 0.15,
        )
        metadata = {
            "from": fake.email(),
            "subject": f"Postulación - {nombre_candidato}",
            "date": fake.date_this_year().isoformat(),
            "attachments": [nombre_archivo],
        }
        (carpeta / "metadata.json").write_text(
            json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        contador += 1

    # --- 2 correos sin adjunto ---
    for _ in range(2):
        email_id = f"email_{contador:03d}"
        carpeta = INBOX_DIR / email_id
        carpeta.mkdir()
        metadata = {
            "from": fake.email(),
            "subject": "Consulta sobre la vacante",
            "date": fake.date_this_year().isoformat(),
            "attachments": [],
        }
        (carpeta / "metadata.json").write_text(
            json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        contador += 1

    # --- 2 correos con formato no soportado ---
    for _ in range(2):
        email_id = f"email_{contador:03d}"
        carpeta = INBOX_DIR / email_id
        carpeta.mkdir()
        nombre_archivo = "foto_cv.jpg"
        (carpeta / nombre_archivo).write_bytes(b"contenido_binario_simulado_de_imagen")
        metadata = {
            "from": fake.email(),
            "subject": "Mi hoja de vida",
            "date": fake.date_this_year().isoformat(),
            "attachments": [nombre_archivo],
        }
        (carpeta / "metadata.json").write_text(
            json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        contador += 1

    # --- 1 correo con .docx vacío (caso borde HU-08) ---
    email_id = f"email_{contador:03d}"
    carpeta = INBOX_DIR / email_id
    carpeta.mkdir()
    nombre_archivo = "cv_vacio.docx"
    Document().save(carpeta / nombre_archivo)  # documento sin párrafos
    metadata = {
        "from": fake.email(),
        "subject": "CV adjunto",
        "date": fake.date_this_year().isoformat(),
        "attachments": [nombre_archivo],
    }
    (carpeta / "metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    contador += 1

    print(f"Bandeja simulada generada en: {INBOX_DIR}")
    print(f"Total de correos creados: {contador - 1}")


if __name__ == "__main__":
    generar_bandeja()
