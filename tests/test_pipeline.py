"""
Pruebas unitarias mínimas para el Sprint 1 (Definition of Done).
Ejecutar con: pytest tests/ -v   (desde la raíz del proyecto)
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document

from src import email_detector, attachment_handler, text_extractor, profile_structurer, schema


def test_deteccion_correo_sin_adjunto(tmp_path):
    carpeta = tmp_path / "email_x"
    carpeta.mkdir()
    (carpeta / "metadata.json").write_text('{"attachments": []}', encoding="utf-8")
    resultado = email_detector.detectar_correo(carpeta)
    assert resultado.estado == "ignorado_sin_adjunto"


def test_deteccion_correo_formato_no_soportado(tmp_path):
    carpeta = tmp_path / "email_y"
    carpeta.mkdir()
    (carpeta / "foto.jpg").write_bytes(b"x")
    (carpeta / "metadata.json").write_text('{"attachments": ["foto.jpg"]}', encoding="utf-8")
    resultado = email_detector.detectar_correo(carpeta)
    assert resultado.estado == "descartado_formato_no_soportado"


def test_deteccion_correo_valido(tmp_path):
    carpeta = tmp_path / "email_z"
    carpeta.mkdir()
    (carpeta / "cv.docx").write_bytes(b"contenido")
    (carpeta / "metadata.json").write_text('{"attachments": ["cv.docx"]}', encoding="utf-8")
    resultado = email_detector.detectar_correo(carpeta)
    assert resultado.estado == "candidato_a_procesar"
    assert resultado.adjunto_valido.name == "cv.docx"


def test_extraccion_adjunto_vacio_es_invalido(tmp_path):
    archivo_vacio = tmp_path / "cv.docx"
    archivo_vacio.write_bytes(b"")
    resultado = attachment_handler.extraer_adjunto("email_1", archivo_vacio, tmp_path / "salida")
    assert resultado.valido is False


def test_extraccion_texto_docx(tmp_path):
    doc = Document()
    doc.add_paragraph("Juan Pérez")
    doc.add_paragraph("juan@example.com")
    doc.add_paragraph("Habilidades")
    doc.add_paragraph("Python, SQL")
    ruta = tmp_path / "cv.docx"
    doc.save(ruta)

    resultado = text_extractor.extraer_texto("email_1", ruta)
    assert resultado.procesable is True
    assert "juan@example.com" in resultado.texto_completo


def test_docx_vacio_no_procesable(tmp_path):
    doc = Document()
    ruta = tmp_path / "cv_vacio.docx"
    doc.save(ruta)
    resultado = text_extractor.extraer_texto("email_1", ruta)
    assert resultado.procesable is False


def test_estructuracion_perfil_no_alucina_campos_faltantes():
    lineas = ["Ana Gómez", "ana@example.com"]
    texto_completo = "\n".join(lineas)
    perfil = profile_structurer.estructurar_perfil("email_1", "cv.docx", texto_completo, lineas)

    assert perfil["nombre"]["valor"] == "Ana Gómez"
    assert perfil["contacto"]["email"]["valor"] == "ana@example.com"
    # No hay sección de educación/experiencia/habilidades en el texto -> deben quedar en None
    assert perfil["educacion"]["valor"] is None
    assert perfil["educacion"]["evidencia"] is None
    assert perfil["experiencia"]["valor"] is None
    assert perfil["habilidades"]["valor"] is None


def test_perfil_valido_contra_schema():
    lineas = ["Carlos Ruiz", "carlos@example.com", "Habilidades", "Python"]
    texto_completo = "\n".join(lineas)
    perfil = profile_structurer.estructurar_perfil("email_1", "cv.docx", texto_completo, lineas)
    es_valido, error = schema.validar_perfil(perfil)
    assert es_valido, error


def test_perfil_invalido_es_rechazado():
    perfil_malformado = {"email_id": "x"}  # faltan campos requeridos
    es_valido, error = schema.validar_perfil(perfil_malformado)
    assert es_valido is False
    assert error is not None


# ============================================================
# Sprint 4 - HU-16
# More robust name heuristic for atypical CV templates
# ============================================================

from src.profile_structurer import _extraer_nombre


def test_hu16_nombre_despues_de_titulo_cv():
    """
    El nombre puede aparecer después de un título genérico del documento.
    """
    lineas = [
        "CURRICULUM VITAE",
        "",
        "Juan Carlos Pérez Gómez",
        "juan.perez@email.com",
        "+57 300 123 4567",
    ]

    nombre, evidencia = _extraer_nombre(lineas)

    assert nombre == "Juan Carlos Pérez Gómez"
    assert evidencia == "Juan Carlos Pérez Gómez"


def test_hu16_nombre_despues_de_datos_contacto():
    """
    El nombre puede aparecer después de información de contacto.
    """
    lineas = [
        "HOJA DE VIDA",
        "Email: maria@email.com",
        "Teléfono: +57 301 555 5555",
        "",
        "María Fernanda López",
        "Experiencia",
    ]

    nombre, evidencia = _extraer_nombre(lineas)

    assert nombre == "María Fernanda López"
    assert evidencia == "María Fernanda López"


def test_hu16_descarta_urls_antes_del_nombre():
    """
    Una URL ubicada antes del nombre no debe ser interpretada
    como nombre del candidato.
    """
    lineas = [
        "RESUME",
        "https://www.linkedin.com/in/candidato",
        "www.portfolio.com",
        "",
        "Carlos Andrés Ramírez",
        "Skills",
    ]

    nombre, evidencia = _extraer_nombre(lineas)

    assert nombre == "Carlos Andrés Ramírez"
    assert evidencia == "Carlos Andrés Ramírez"


def test_hu16_nombre_fuera_de_primeras_tres_lineas():
    """
    La heurística debe encontrar nombres que aparecen después
    de las primeras tres líneas del CV.
    """
    lineas = [
        "CURRICULUM VITAE",
        "",
        "email: ana@email.com",
        "+57 310 555 1234",
        "",
        "Ana Sofía Martínez",
        "Educación",
    ]

    nombre, evidencia = _extraer_nombre(lineas)

    assert nombre == "Ana Sofía Martínez"
    assert evidencia == "Ana Sofía Martínez"


def test_hu16_no_inventa_nombre_sin_candidato_valido():
    """
    Si no existe evidencia suficiente para identificar un nombre,
    la función debe mantener el principio anti-alucinación.
    """
    lineas = [
        "CURRICULUM VITAE",
        "email: candidato@email.com",
        "+57 300 123 4567",
        "https://www.linkedin.com/in/candidato",
        "Experiencia",
    ]

    nombre, evidencia = _extraer_nombre(lineas)

    assert nombre is None
    assert evidencia is None